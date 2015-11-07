from django.shortcuts import render
from django.conf import settings
from django.db import IntegrityError

from .models import Guest
from register.models import GuestRegistration
from branchoffice.models import BranchOffice

import datetime
import openpyxl

def index(request):
    return render(request, 'guest.html')

def cleanGuestDocument(request):
    guests = Guest.objects.all()
    for guest in guests:
        try:
            vd = guest.val_document
            val = vd.split('.')[0]
            guest.val_document = val
            guest.save()
        except IntegrityError:
            guest.delete()

def importGuest(request):
    workbook = openpyxl.load_workbook(filename = settings.MEDIA_ROOT + '/'
                                      + settings.UPLOAD_PATH
                                      + '/files/registro_personas.xlsx',
                                      use_iterators = True)

    worksheet = workbook.get_sheet_by_name('registros')

    for row in worksheet.iter_rows():
        data = {
                'ci': str(row[0].value).split('.')[0],
                'nombre': row[1].value,
                'apellido': row[2].value,
                'ciudad': row[3].value,
                'sexo': row[4].value,
                'fecha': datetime.datetime.strptime((row[5].value), '%d/%m/%Y').date(),
                'entrada': row[6].value,
                'salida': row[7].value,
                'motivo': row[8].value,
        }
        loadInDBregister(request, data)
    print 'done.html'

def loadInDBregister(request, datalist):
    try:
        guest = Guest.objects.get(val_document=datalist['ci'])
    except Guest.DoesNotExist:
        guest = Guest(first_name=datalist['nombre'],
                    last_name=datalist['apellido'],
                    date_created=datalist['fecha'],
                    type_document='CI',
                    issue_document=datalist['ciudad'],
                    val_document=datalist['ci'],
                    gender='H' if datalist['sexo'] == 'M' else 'M')
        guest.save()

    muyurina = BranchOffice.objects.get(name='Muyurina')
    gr = GuestRegistration(guest=guest,
                           register_date=str(datalist['fecha']),
                           time_entry=str(datalist['fecha']) + ' ' + datalist['entrada'],
                           time_out=str(datalist['fecha']) + ' ' + datalist['salida'],
                           reason=datalist['motivo'],
                           branchoffice=muyurina,
                           owner=request.user)
    gr.save()

